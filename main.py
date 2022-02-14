import os
import datetime
from math import isnan

import openpyxl
import pandas as pd
from docx import Document
from docx2pdf import convert
from openpyxl import Workbook
from xlwt import Workbook as Wbk


def make_letter(dic, sheet_name, ind):
    document = Document('letter.docx')
    todays_date = datetime.date.today()
    todays_date = todays_date.strftime("%d.%m.%Y")
    for p in document.paragraphs:
        inline = p.runs
        for i in range(len(inline)):
            text = inline[i].text
            for key in dic.keys():
                if key in text:
                    text = text.replace(key, dic[key])
                    inline[i].text = text
    if os.path.exists('wygenerowane') == False:
        os.mkdir('wygenerowane', 0o666)
    if os.path.exists('wygenerowane/' + todays_date) == False:
        os.mkdir('wygenerowane/' + todays_date, 0o666)
    document.save('wygenerowane/' + todays_date + '/' + sheet_name + '_' + ind + '.docx')
    convert('wygenerowane/' + todays_date + '/' + sheet_name + '_' + ind + '.docx',
            'wygenerowane/' + todays_date + '/' + sheet_name + '_' + ind + '.pdf')
    os.remove('wygenerowane/' + todays_date + '/' + sheet_name + '_' + ind + '.docx')

def makeXlsForPostoffice(addresseeList_):
    do_importu_file = pd.ExcelFile('doimportu.xls')
    sheets = do_importu_file.sheet_names
    filename = datetime.datetime.now(datetime.timezone.utc).strftime("%d-%m-%Y %H %M")
    filepath = 'wygenerowane/import/' + filename +'.xls'
    if os.path.exists('wygenerowane') == False:
        os.mkdir('wygenerowane', 0o666)
    if os.path.exists('wygenerowane/import') == False:
        os.mkdir('wygenerowane/import', 0o666)



    book = Wbk(filepath)

    book.add_sheet(sheets[0])
    book.save(filepath)
    to_import = pd.read_excel('doimportu.xls', sheets[0])
    for i in addresseeList_:
        to_import.loc[addresseeList_.index(i), "AdresatNazwa"] = i[0]
        to_import.loc[addresseeList_.index(i), "AdresatUlica"] = i[1]
        to_import.loc[addresseeList_.index(i), "AdresatNumerDomu"] = i[2]
        to_import.loc[addresseeList_.index(i), "AdresatNumerLokalu"] = i[3]
        to_import.loc[addresseeList_.index(i), "AdresatKodPocztowy"] = i[4]
        to_import.loc[addresseeList_.index(i), "AdresatMiejscowosc"] = i[5]
        to_import.loc[addresseeList_.index(i), "AdresatKraj"] = i[6]
        to_import.loc[addresseeList_.index(i), "Format"] = "S"
        to_import.loc[addresseeList_.index(i), "KategoriaLubGwarancjaTerminu"] = "E"



    with pd.ExcelWriter(filepath, mode='w', engine='xlwt') as writer:
            to_import.to_excel(writer, sheet_name=sheets[0], index=False)


def join_name(matrix):
    name = ''
    for i in matrix:
        for j in i:
            name += str(j)
        name += ' '
    return name


def to_uppercase(name):
    matrix_temp = name.split(" ")
    matrix = []
    for i in matrix_temp:
        matrix.append(list(i))
    for i in range(len(matrix)):
        for j in range(len(matrix[i])):
            if j == 0:
                matrix[i][j] = str(matrix[i][j]).upper()
            else:
                matrix[i][j] = str(matrix[i][j]).lower()
    name_ = join_name(matrix)
    return name_

def reverseName(name):
    name = name.split()
    name = name[::-1]
    return join_name(name)




def main():
    todays_date = datetime.date.today()
    todays_date = todays_date.strftime("%d.%m.%Y")
    iterator = 0
    xls = pd.ExcelFile('dane.xlsx')
    sheet_nam = xls.sheet_names

    count = int(input("How many letters you want to generate: \n"))
    book = Workbook("dane_kopia.xlsx")
    book.save("dane_kopia.xlsx")
    book.close()

    addresseeList = [] # [[NazwiskoImie, Ulica, Nrdomu, NrLok, kodpoczt, Miejscowosc, 'Polska'], ...]

    for i in sheet_nam:
        df = pd.read_excel('dane.xlsx', i)
        for j in df.itertuples():
            if pd.isna(j.nazwa) or pd.isna(j._3) or pd.isna(j._4) or pd.isna(j._5) or pd.isna(
                    j._6) or pd.isna(j._12) or pd.isna(j._14):
                continue
            if (j.generated == 'nie') and iterator < count:
                if type(j._7) == str:
                    nr_lok = j._7
                else:
                    if isnan(j._7):
                        nr_lok = ''
                    else:
                        nr_lok = '/' + str(j._7)
                kod =str(j._3)
                if len(kod) == 6:
                    kod = kod
                else:
                    kod = str(int(j._3))
                    kod = kod[0] + kod[1] + '-' + kod[2] + kod[3] + kod[4]

                data = j._12.strftime("%d.%m.%Y")

                imie_nazwisko = to_uppercase(j.nazwa)
                nazwa_ulicy = to_uppercase(j._5)
                miejscowosc = to_uppercase(j._4)
                kwota_dodatkowa_tab = j._15.split(" ")

                kwota_dodatkowa_tab[0] = "{:.2f}".format(float(kwota_dodatkowa_tab[0]) + 10)
                kwota_dodatkowa = join_name(kwota_dodatkowa_tab)

                dic = {
                    '1@': imie_nazwisko,  # imie i nazwisko
                    '1!': str(todays_date),  # dzisiejsza data
                    '2!': nazwa_ulicy,  # nazwa ulicy
                    '2@': str(j._6),  # numer domu
                    '2#': nr_lok,  # numer lokalu
                    '3!': kod,  # kod pocztowy
                    '3@': miejscowosc,  # miejscowosc
                    '4!': str(j._15),  # podstawa roszczenia
                    '4@': str(data),  # data roszczenia
                    '5!': kwota_dodatkowa,  # roszczenie +10
                    'ind': str(j.Index + 1),
                }
                make_letter(dic, str(i), str(j.Index + 1))
                df.loc[j.Index, "generated"] = 'tak'
                addresseeList.append((reverseName(imie_nazwisko), nazwa_ulicy, str(j._6), nr_lok[1:], join_name(kod.split('-')), miejscowosc, 'Polska'))
                iterator += 1

        with pd.ExcelWriter("dane_kopia.xlsx", mode='a') as writer:
            df.to_excel(writer, sheet_name=i, index=False)

    workbook = openpyxl.load_workbook("dane_kopia.xlsx")
    stf = workbook['Sheet']
    workbook.remove(stf)
    workbook.save("dane_kopia.xlsx")
    makeXlsForPostoffice(addresseeList)
    xls.close()
    os.remove('dane.xlsx')
    os.rename("dane_kopia.xlsx", "dane.xlsx")
